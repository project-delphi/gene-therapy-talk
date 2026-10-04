"""Render the slide model into a .pptx, on top of a themed base deck.

Everything is positioned by hand rather than through pandoc's pptx writer,
which breaks this deck: it starts a new slide for anything following a
`.columns` block, orphaning every citation and every stacked figure.

Geometry is expressed as fractions of the slide box so one renderer serves
both page sizes (ACOI is 13.333x7.5in, the original deck 10x5.625in).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import ClassVar

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckmodel import Bullets, Figure, Slide, Text

FONTS = {
    (False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
    (True, False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    (False, True): "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
    (True, True): "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf",
}
MONO_FONT = "/System/Library/Fonts/Supplemental/Courier New.ttf"
_REF_PX = 200
_font_cache: dict = {}


def _font(bold: bool, italic: bool, mono: bool):
    key = (bold, italic, mono)
    if key not in _font_cache:
        path = MONO_FONT if mono else FONTS[(bold, italic)]
        _font_cache[key] = ImageFont.truetype(path, _REF_PX)
    return _font_cache[key]


def text_width(text: str, size_pt: float, bold=False, italic=False, mono=False) -> float:
    """Width of `text` in points when set at `size_pt`."""
    return _font(bold, italic, mono).getlength(text) / _REF_PX * size_pt


# --- theme -----------------------------------------------------------------


@dataclass
class Theme:
    name: str
    base: Path
    heading: str
    body: str
    bang: str
    link: str
    accent: str
    cite: str
    logo: Path | None = None
    subtitle: list[str] = field(default_factory=list)
    cite_top: float = 0.905
    cite_size: float = 10.0     # points, quoted against a 13.333in-wide slide
    logo_box: tuple[float, float, float] | None = None  # x, y, size (fractions)

    def body_bottom_at(self, x: float) -> float:
        """Lowest y (fraction) content may reach at horizontal position `x`."""
        return 0.885


# Point sizes below are quoted against a 13.333in-wide slide and scaled for the
# narrower original deck, so both come out the same visual size.
REF_SLIDE_IN = 13.3333
BASE_BODY_PT = 20.0
MIN_BODY_PT = 11.0
TITLE_PT = 30.0
SECTION_PT = 46.0
COVER_PT = 54.0

MARGIN_L = 0.055
MARGIN_R = 0.055
TITLE_TOP = 0.045
BODY_TOP = 0.170
COL_GAP = 0.022
ROW_GAP = 0.022
TAIL_SPAN = 0.64   # width of a trailing one-liner, keeping it off the artwork


class AcoiTheme(Theme):
    # Top edge of the green swoosh, measured off the master's background art
    # by sampling where the white ends in each column. Content stays above it.
    CURVE: ClassVar[list[tuple[float, float]]] = [(0.00, 0.950), (0.10, 0.921), (0.20, 0.910), (0.30, 0.896),
             (0.40, 0.880), (0.50, 0.859), (0.60, 0.836), (0.70, 0.807),
             (0.80, 0.772), (0.90, 0.727), (1.00, 0.655)]

    def body_bottom_at(self, x: float) -> float:
        x = min(max(x, 0.0), 1.0)
        for (x0, y0), (x1, y1) in zip(self.CURVE, self.CURVE[1:]):
            if x <= x1:
                t = (x - x0) / (x1 - x0)
                return min(y0 + t * (y1 - y0) - 0.018, 0.800)
        return 0.800


def acoi_theme(repo: Path, template: Path) -> Theme:
    return AcoiTheme(
        name="acoi",
        base=template,
        heading="47744B", body="1A1A1A", bang="47744B",
        link="2F6E7A", accent="99C160", cite="63706A",
        logo=repo / "tools" / "assets" / "acoi-logo.png",
        subtitle=["Ravi Kalia, PhD",
                  "VIII Encuentro de Investigación en Inmunología",
                  "Asociación Colombiana de Inmunología · Universidad del Magdalena · 2026"],
        cite_top=0.815,
        cite_size=10.0,
        logo_box=(0.0081, 0.0313, 0.1338),
    )


def original_theme(repo: Path) -> Theme:
    return Theme(
        name="original",
        base=repo / "source" / "gene_therapy.pptx",
        heading="FFFFFF", body="ADADAD", bang="26A69A",
        link="4DD0E1", accent="009688", cite="8E8E8E",
        subtitle=["Ravi Kalia"],
        cite_top=0.905,
        cite_size=10.0,
    )


# --- presentation helpers --------------------------------------------------


def clear_slides(prs) -> None:
    id_list = prs.slides._sldIdLst
    for sld in list(id_list):
        prs.part.drop_rel(sld.rId)
        id_list.remove(sld)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.element.get("type") == "blank":
            return layout
    return prs.slide_layouts[-1]


class Canvas:
    """Fraction-to-EMU helper bound to one presentation's page size."""

    def __init__(self, prs):
        self.w = prs.slide_width
        self.h = prs.slide_height
        # Everything authored for a 13.333in slide shrinks proportionally.
        self.scale = (self.w / 914400) / REF_SLIDE_IN

    def x(self, f): return Emu(int(self.w * f))
    def y(self, f): return Emu(int(self.h * f))
    def pt(self, p): return Pt(p * self.scale)

    def pt_to_frac(self, p: float) -> float:
        """Height of `p` points as a fraction of slide height."""
        return (p * 12700) / self.h

    def frac_to_pt(self, f: float) -> float:
        return f * self.h / 12700


# --- text layout -----------------------------------------------------------

BULLET_INDENT = [0.0, 0.30]   # inches per level, before scaling
BULLET_HANG = 0.26            # inches reserved for the glyph
LINE_SPACING = 1.20
PARA_GAP = 0.45               # multiples of font size, between blocks
BULLET_GAP = 0.16             # multiples of font size, between bullets


@dataclass
class Para:
    runs: list
    level: int | None      # None => not a bullet
    space_before: float    # multiples of font size
    space_after: float
    italic: bool = False


def plan(blocks) -> list[Para]:
    """Flatten a column's text blocks into the paragraphs we will emit.

    Measurement and rendering both read this, so the height estimate and the
    shape that lands on the slide cannot drift apart.
    """
    paras: list[Para] = []
    for i, blk in enumerate(blocks):
        gap = PARA_GAP if i else 0.0
        if isinstance(blk, Bullets):
            for j, (level, runs) in enumerate(blk.items):
                paras.append(Para(runs, level, gap if j == 0 else 0.0, BULLET_GAP))
        elif isinstance(blk, Text):
            paras.append(Para(blk.runs, None, gap, 0.0, italic=blk.quote))
    return paras


def _avail_pt(para: Para, width_in: float, scale: float) -> float:
    if para.level is None:
        return width_in * 72
    indent = (BULLET_INDENT[min(para.level, 1)] + BULLET_HANG) * scale
    return (width_in - indent) * 72


def _wrap_count(runs, size_pt: float, width_pt: float) -> int:
    if width_pt <= 0:
        return 1
    words: list[tuple[str, object]] = []
    for run in runs:
        parts = run.text.split(" ")
        for i, part in enumerate(parts):
            if i:
                words.append((" ", run))
            if part:
                words.append((part, run))
    lines, cur = 1, 0.0
    for word, run in words:
        w = text_width(word, size_pt, run.bold, run.italic, run.mono)
        if word == " ":
            cur += w
            continue
        if cur + w > width_pt and cur > 0:
            lines += 1
            cur = w
        else:
            cur += w
    return lines


def text_height_pt(paras: list[Para], size_pt: float, width_in: float,
                   scale: float) -> float:
    total = 0.0
    for para in paras:
        lines = _wrap_count(para.runs, size_pt, _avail_pt(para, width_in, scale))
        total += (para.space_before * size_pt
                  + lines * size_pt * LINE_SPACING
                  + para.space_after * size_pt)
    return total


def fit_size(paras, width_in, scale, budget_pt, base_pt) -> float:
    size = base_pt
    while size > MIN_BODY_PT and \
            text_height_pt(paras, size * scale, width_in, scale) > budget_pt:
        size -= 0.5
    return size


_ROLE_COLOUR = {"body": "body", "bang": "bang", "link": "link", "quote": "cite"}


def write_paras(tf, paras: list[Para], theme: Theme, size_pt: float,
                scale: float) -> None:
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(size_pt * LINE_SPACING)
        p.space_before = Pt(para.space_before * size_pt)
        p.space_after = Pt(para.space_after * size_pt)
        if para.level is not None:
            _bullet(p, para.level, scale, theme)
        for r in para.runs:
            run = p.add_run()
            run.text = r.text
            f = run.font
            f.size = Pt(size_pt)
            f.name = "Courier New" if r.mono else "Arial"
            f.bold = r.bold
            f.italic = r.italic or para.italic
            f.color.rgb = RGBColor.from_string(getattr(theme, _ROLE_COLOUR[r.role]))
            if r.href:
                run.hyperlink.address = r.href
                f.color.rgb = RGBColor.from_string(theme.link)


def _bullet(p, level: int, scale: float, theme: Theme) -> None:
    p.level = min(level, 1)
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int((BULLET_INDENT[min(level, 1)] + BULLET_HANG)
                            * scale * 914400)))
    pPr.set("indent", str(-int(BULLET_HANG * scale * 914400)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": theme.body}))
    pPr.append(clr)
    # Matches the disc / circle glyphs the reveal theme uses for the two levels.
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "◦"}))


# --- figure layout ---------------------------------------------------------


def _fit(size, box_w, box_h):
    iw, ih = size
    s = min(box_w / iw, box_h / ih)
    return int(iw * s), int(ih * s)


def _stack(figs, sizes, w, h, gap, spot_h):
    """Size a vertical stack of figures into a `w` x `h` box.

    Mirrors the shrinking per-image height cap in theme.scss: the more figures
    share a column, the less height each one may claim.
    """
    spots = sum(1 for f in figs if f.spot)
    main = len(figs) - spots
    room = h - gap * (len(figs) - 1) - spot_h * spots
    each = max(room / max(main, 1), 1)
    out = []
    for fig, size in zip(figs, sizes):
        out.append(_fit(size, w, spot_h if fig.spot else each))
    return out


def place_figures(slide, repo, canvas, theme, figs, x, y, w, h_max,
                  clearance: bool):
    """Stack figures in the column, shrinking until they clear the background.

    ACOI's swoosh rises from the bottom-right, so the usable height depends on
    how far right the artwork actually reaches. Sizing against the column's
    right edge wastes a lot of room, because a centred figure is usually
    narrower than its column — so size generously, then shrink only if the
    placed pixels really do collide.
    """
    sizes = []
    for fig in figs:
        with Image.open(repo / fig.path) as im:
            sizes.append(im.size)

    gap = int(canvas.h * 0.012)
    spot_h = int(canvas.h * 0.105)
    h = h_max
    avail = h_max
    dims = _stack(figs, sizes, w, h, gap, spot_h)
    for _ in range(10):
        dims = _stack(figs, sizes, w, h, gap, spot_h)
        total = sum(d[1] for d in dims) + gap * (len(dims) - 1)
        if not clearance:
            avail = h_max
            break
        # The lowest figure decides: a centred figure is usually narrower than
        # its column, so it clears the swoosh further down than the column edge
        # would suggest.
        pw, ph = dims[-1]
        right = (x + (w - pw) / 2 + pw) / canvas.w
        avail = min(y + h_max, theme.body_bottom_at(right) * canvas.h) - y
        if total <= avail:
            break
        h = int(h * max(avail / total, 0.4))

    total = sum(d[1] for d in dims) + gap * (len(dims) - 1)
    cursor = y + max(0, (avail - total) // 2)
    for fig, (pw, ph) in zip(figs, dims):
        pic = slide.shapes.add_picture(
            str(repo / fig.path), Emu(int(x + (w - pw) / 2)), Emu(int(cursor)),
            Emu(pw), Emu(ph))
        if fig.href:
            pic.click_action.hyperlink.address = fig.href
        cursor += ph + gap


# --- slide rendering -------------------------------------------------------


def _add_title(slide, canvas, theme, title, size_pt, top, height, centre=False):
    box = slide.shapes.add_textbox(
        canvas.x(MARGIN_L), canvas.y(top),
        canvas.x(1 - MARGIN_L - MARGIN_R), canvas.y(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if centre else MSO_ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER if centre else PP_ALIGN.LEFT
    avail = (1 - MARGIN_L - MARGIN_R) * canvas.w / 914400 * 72
    while size_pt > 14 and text_width(title, size_pt * canvas.scale) > avail:
        size_pt -= 1
    run = para.add_run()
    run.text = title
    run.font.size = canvas.pt(size_pt)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(theme.heading)
    return box


def _accent_rule(slide, canvas, theme, top, width_frac=0.16):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        canvas.x(0.5 - width_frac / 2), canvas.y(top),
        canvas.x(width_frac), Emu(max(int(canvas.h * 0.004), 12700)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(theme.accent)
    bar.line.fill.background()
    bar.shadow.inherit = False


def _render_column(slide, repo, canvas, theme, col, x, y, w, h, base_pt,
                   clearance: bool):
    texts = [b for b in col.blocks if isinstance(b, (Bullets, Text))]
    figs = [b for b in col.blocks if isinstance(b, Figure)]
    width_in = w / 914400
    avail_pt = canvas.frac_to_pt(h / canvas.h)
    base = base_pt * (0.80 if "tight" in col.classes else 1.0)

    if texts:
        paras = plan(texts)
        # A column holding both text and figures caps the text at two thirds so
        # the figure is not squeezed to nothing.
        budget = avail_pt * (0.66 if figs else 1.0)
        size = fit_size(paras, width_in, canvas.scale, budget, base)
        used = text_height_pt(paras, size * canvas.scale, width_in, canvas.scale)
        used = min(used, avail_pt)
        box = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)),
                                       Emu(int(used * 12700)))
        write_paras(box.text_frame, paras, theme, size * canvas.scale, canvas.scale)
        if figs:
            consumed = int((used + size * canvas.scale * PARA_GAP) * 12700)
            y += consumed
            h -= consumed

    if figs and h > 0:
        place_figures(slide, repo, canvas, theme, figs, x, y, w, h, clearance)


def _render_cite(slide, canvas, theme, runs):
    box = slide.shapes.add_textbox(
        canvas.x(MARGIN_L), canvas.y(theme.cite_top),
        canvas.x(0.46), canvas.y(0.055))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for r in runs:
        run = tf.paragraphs[0].add_run()
        run.text = r.text
        run.font.size = canvas.pt(theme.cite_size)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(theme.cite)
        if r.href:
            run.hyperlink.address = r.href


def _render_content(slide, repo, canvas, theme, s: Slide):
    _add_title(slide, canvas, theme, s.title, TITLE_PT, TITLE_TOP, 0.13)

    span = 1 - MARGIN_L - MARGIN_R
    floor = theme.cite_top - 0.030 if s.cite else theme.body_bottom_at(1.0)

    # Rows after the first are the shouted one-liners that follow a columns
    # block. They sit just above the citation; the first row takes the rest.
    tail = s.rows[1:]
    tail_paras = []
    tail_h = 0.0
    # Held back from the full measure so the line cannot run out over the
    # swoosh on the right; it wraps instead.
    tail_span = min(span, TAIL_SPAN)
    tail_in = tail_span * canvas.w / 914400
    for row in tail:
        paras = plan([b for c in row.columns for b in c.blocks
                      if isinstance(b, (Bullets, Text))])
        size = fit_size(paras, tail_in, canvas.scale,
                        canvas.frac_to_pt(0.16), BASE_BODY_PT)
        hgt = canvas.pt_to_frac(
            text_height_pt(paras, size * canvas.scale, tail_in, canvas.scale))
        tail_paras.append((paras, size, hgt))
        tail_h += hgt + ROW_GAP

    first_bottom = floor - tail_h if tail else floor

    if s.rows:
        row = s.rows[0]
        total = sum(c.width for c in row.columns) or 1.0
        usable = span - COL_GAP * (len(row.columns) - 1)
        x = MARGIN_L
        for col in row.columns:
            cw = usable * col.width / total
            # Text has to clear the artwork on its own; figures negotiate it in
            # place_figures, which knows how wide they actually are.
            has_text = any(isinstance(b, (Bullets, Text)) for b in col.blocks)
            limit = min(first_bottom, theme.body_bottom_at(x + cw)) if has_text \
                else max(first_bottom, theme.body_bottom_at(1.0))
            _render_column(slide, repo, canvas, theme, col,
                           canvas.x(x), canvas.y(BODY_TOP), canvas.x(cw),
                           canvas.y(max(limit - BODY_TOP, 0.06)),
                           BASE_BODY_PT, clearance=not has_text)
            x += cw + COL_GAP

    y = floor - tail_h + ROW_GAP / 2
    for paras, size, hgt in tail_paras:
        box = slide.shapes.add_textbox(
            canvas.x(MARGIN_L), canvas.y(y), canvas.x(tail_span),
            canvas.y(max(hgt, 0.04)))
        write_paras(box.text_frame, paras, theme, size * canvas.scale, canvas.scale)
        y += hgt + ROW_GAP

    if s.cite:
        _render_cite(slide, canvas, theme, s.cite)


def render(meta, slides: list[Slide], theme: Theme, repo: Path, out: Path,
           only: list[str] | None = None) -> int:
    prs = Presentation(str(theme.base))
    clear_slides(prs)
    layout = blank_layout(prs)
    canvas = Canvas(prs)

    if only is not None:
        keep = set(only)
        missing = keep - {s.title for s in slides}
        if missing:
            # Headings get rewritten. Exiting here took the other two decks and
            # every PDF down with the short cut, and made the caller's softer
            # "the 10-minute cut will be short" warning a lie. Build what is
            # still there and say what was dropped.
            print(f"  !  short deck names {len(missing)} slide(s) that no longer exist: "
                  f"{', '.join(repr(t) for t in sorted(missing))}")
        slides = [s for s in slides if s.title in keep]

    # --- title slide
    slide = prs.slides.add_slide(layout)
    if theme.logo and theme.logo_box:
        lx, ly, ls = theme.logo_box
        side = int(canvas.w * ls)
        slide.shapes.add_picture(str(theme.logo), canvas.x(lx), canvas.y(ly),
                                 Emu(side), Emu(side))
    _add_title(slide, canvas, theme, meta.get("title", "Presentation"),
               COVER_PT, 0.32, 0.18, centre=True)
    _accent_rule(slide, canvas, theme, 0.545, 0.20)
    sub = slide.shapes.add_textbox(canvas.x(MARGIN_L), canvas.y(0.605),
                                   canvas.x(1 - MARGIN_L - MARGIN_R), canvas.y(0.20))
    stf = sub.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
    for i, line in enumerate(theme.subtitle):
        para = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        para.line_spacing = 1.4
        run = para.add_run()
        run.text = line
        run.font.size = canvas.pt(23 if i == 0 else 16)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(theme.body)

    for s in slides:
        slide = prs.slides.add_slide(layout)
        if s.kind == "section":
            _add_title(slide, canvas, theme, s.title, SECTION_PT, 0.35, 0.16,
                       centre=True)
            _accent_rule(slide, canvas, theme, 0.545)
        else:
            _render_content(slide, repo, canvas, theme, s)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return len(prs.slides._sldIdLst)
