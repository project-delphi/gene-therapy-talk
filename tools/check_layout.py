#!/usr/bin/env python3
"""Geometry check for the generated decks.

Reads each shape's real bounding box out of the .pptx and flags anything that
runs off the slide or, for the ACOI deck, collides with the green swoosh on the
master background. Cheaper and more reliable than eyeballing 38 slides.

    uv run --with python-pptx python3 tools/check_layout.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckrender import AcoiTheme  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "viii-encuentro-acoi-unimagdalena"


def check(path: Path) -> list[str]:
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    acoi = "acoi" in path.stem
    curve = AcoiTheme.body_bottom_at.__get__(AcoiTheme("x", path, *["0"] * 6))
    problems = []

    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.width is None or shape.height is None:
                continue
            left, top = shape.left / W, shape.top / H
            right, bottom = (shape.left + shape.width) / W, (shape.top + shape.height) / H
            name = shape.shape_type
            if left < -0.001 or top < -0.001 or right > 1.001 or bottom > 1.001:
                problems.append(
                    f"{path.name} slide {n}: {name} off-slide "
                    f"[{left:.3f},{top:.3f} .. {right:.3f},{bottom:.3f}]")
            # Only pictures are checked against the artwork: text boxes are
            # sized to their measured text and the citation deliberately sits
            # low on the left, where the swoosh has not yet risen.
            if acoi and str(name).startswith("PICTURE") and n > 1:
                limit = curve(right)
                if bottom > limit + 0.004:
                    problems.append(
                        f"{path.name} slide {n}: picture bottom {bottom:.3f} "
                        f"past swoosh limit {limit:.3f} at x={right:.3f}")
    return problems


def main() -> None:
    found = 0
    for pptx in sorted(OUT.glob("*.pptx")):
        problems = check(pptx)
        found += len(problems)
        print(f"{pptx.name}: {len(problems)} problem(s)")
        for p in problems:
            print("   ", p)
    sys.exit(1 if found else 0)


if __name__ == "__main__":
    main()
